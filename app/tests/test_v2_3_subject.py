"""V2.3 测试：公司维度（仓库所属公司 / 单据签约公司 / 调拨单位 / 库存按公司筛选）。

对应本轮需求（`docs/26-v2.3-adjustments.md`）：
- N1 仓库新增「所属公司」（只能取系统设置的我方公司，主体码），其他表单仓库默认显示 `ZC-主仓`；
- N2 采购申请单 / 采购单 / 销售申请单 / 销售订单 / 入库单 / 出库单 新增「签约公司」（必填）；
- N3 库存明细按公司筛选商品（口径 = 仓库所属公司）；
- N4 调拨单的调出单位 / 调入单位由两端仓库所属公司自动带出（只读快照）。

测试自建主数据与单据，结束时按依赖顺序清理（不留脏数据）。
"""
from __future__ import annotations

import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.database import SessionLocal
from app.init_db import ADMIN_INIT_PASSWORD, ADMIN_USERNAME
from app.main import app
from app.models import ChangeLog
from app.models_auth import User

SUBJECT = "ZC"           # 系统设置默认主体（智澈公司）
SUBJECT_NAME = "智澈公司"
SUBJECT2 = "YX"          # 第二个主体（云羲公司）

DOC_CREATE_PATHS = [
    "/api/purchase/requests",
    "/api/purchase/orders",
    "/api/sales/requests",
    "/api/sales/orders",
    "/api/stock/in-orders",
    "/api/stock/out-orders",
]

# 接口前缀 → (行项表, 主表, 变更历史的 object_type)
DOC_TABLES: dict[str, tuple[str, str, str]] = {
    "/api/purchase/requests": ("purchase_request_items", "purchase_requests", "purchase_request"),
    "/api/purchase/orders": ("purchase_order_items", "purchase_orders", "purchase_order"),
    "/api/sales/requests": ("sales_request_items", "sales_requests", "sales_request"),
    "/api/sales/orders": ("sales_order_items", "sales_orders", "sales_order"),
    "/api/stock/in-orders": ("stock_in_order_items", "stock_in_orders", "stock_in"),
    "/api/stock/out-orders": ("stock_out_order_items", "stock_out_orders", "stock_out"),
    "/api/stock/transfers": ("stock_transfer_items", "stock_transfers", "stock_transfer"),
}


def _sfx() -> str:
    return uuid.uuid4().hex[:6]


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def admin(client):
    resp = client.post("/api/auth/login",
                       json={"username": ADMIN_USERNAME, "password": ADMIN_INIT_PASSWORD})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['token']}"}


@pytest.fixture()
def db():
    with SessionLocal() as s:
        yield s


class Factory:
    """测试数据工厂（含清理）。"""

    def __init__(self, client: TestClient, admin: dict) -> None:
        self.client = client
        self.admin = admin
        self.warehouses: list[int] = []
        self.docs: dict[str, list[int]] = {}
        self.products: list[int] = []
        self.uoms: list[int] = []
        self.types: list[int] = []
        self.suppliers: list[int] = []
        self.customers: list[int] = []

    # ---- 主数据 ----
    def uom(self) -> dict:
        r = self.client.post("/api/master/uoms", headers=self.admin,
                             json={"code": f"U{_sfx()}", "name": f"单位{_sfx()}", "decimals": 2})
        assert r.status_code == 200, r.text
        self.uoms.append(r.json()["id"])
        return r.json()

    def ptype(self) -> dict:
        r = self.client.post("/api/master/product-types", headers=self.admin,
                             json={"name": f"类型{_sfx()}"})
        assert r.status_code == 200, r.text
        self.types.append(r.json()["id"])
        return r.json()

    def product(self) -> dict:
        r = self.client.post("/api/master/products", headers=self.admin, json={
            "name": f"物料{_sfx()}", "product_type_id": self.ptype()["id"],
            "uom_id": self.uom()["id"], "default_price": 10})
        assert r.status_code == 200, r.text
        self.products.append(r.json()["id"])
        return r.json()

    def warehouse(self, *, subject: str | None = SUBJECT) -> dict:
        payload = {"code": f"W{_sfx().upper()}", "name": f"仓{_sfx()}"}
        if subject is not None:
            payload["subject_code"] = subject
        r = self.client.post("/api/master/warehouses", headers=self.admin, json=payload)
        assert r.status_code == 200, r.text
        self.warehouses.append(r.json()["id"])
        return r.json()

    def supplier(self) -> dict:
        r = self.client.post("/api/master/suppliers", headers=self.admin,
                             json={"name": f"供应商{_sfx()}", "short_name": f"供{_sfx()}"})
        assert r.status_code == 200, r.text
        self.suppliers.append(r.json()["id"])
        return r.json()

    def customer(self) -> dict:
        r = self.client.post("/api/master/customers", headers=self.admin,
                             json={"name": f"客户{_sfx()}"})
        assert r.status_code == 200, r.text
        self.customers.append(r.json()["id"])
        return r.json()

    # ---- 单据 ----
    def create_doc(self, path: str, payload: dict, *, expect: int = 200) -> dict:
        r = self.client.post(path, headers=self.admin, json=payload)
        assert r.status_code == expect, r.text
        if expect != 200:
            return {"_error": r.json()}
        doc = r.json()
        self.docs.setdefault(DOC_TABLES[path][1], []).append(doc["id"])
        return doc

    def cleanup(self) -> None:
        with SessionLocal() as db:
            for main_table, ids in self.docs.items():
                item_table, _, object_type = next(
                    t for t in DOC_TABLES.values() if t[1] == main_table)
                for doc_id in ids:
                    db.execute(text(f"DELETE FROM {item_table} WHERE doc_id = :i"), {"i": doc_id})
                    db.execute(text(f"DELETE FROM {main_table} WHERE id = :i"), {"i": doc_id})
                db.query(ChangeLog).filter(ChangeLog.object_type == object_type).delete(
                    synchronize_session=False)
            for pid in self.products:
                db.execute(text("DELETE FROM stock_ledger WHERE product_id = :i"), {"i": pid})
                db.execute(text("DELETE FROM stocks WHERE product_id = :i"), {"i": pid})
            for wid in self.warehouses:
                db.execute(text("DELETE FROM stock_ledger WHERE warehouse_id = :i"), {"i": wid})
                db.execute(text("DELETE FROM stocks WHERE warehouse_id = :i"), {"i": wid})
            for table, ids in (("products", self.products), ("uoms", self.uoms),
                               ("product_types", self.types), ("warehouses", self.warehouses),
                               ("suppliers", self.suppliers), ("customers", self.customers)):
                for obj_id in ids:
                    db.execute(text(f"DELETE FROM {table} WHERE id = :i"), {"i": obj_id})
            db.commit()


@pytest.fixture()
def box(client, admin):
    f = Factory(client, admin)
    try:
        yield f
    finally:
        f.cleanup()


def _items(product: dict, qty: float = 1, price: float = 10) -> list[dict]:
    return [{"product_id": product["id"], "qty": qty, "unit_price": price}]


# ===================== N1 仓库所属公司 =====================

class TestWarehouseSubject:
    def test_所属公司可设置并返回显示名(self, box):
        wh = box.warehouse(subject=SUBJECT)
        assert wh["subject_code"] == SUBJECT
        assert wh["subject_name"] == SUBJECT_NAME
        assert wh["display_name"] == f"{SUBJECT}-{wh['name']}", "显示名必须是 主体码-仓库名"

    def test_所属公司只能是字典内的主体(self, box):
        r = box.client.post("/api/master/warehouses", headers=box.admin,
                            json={"code": f"W{_sfx().upper()}", "name": f"仓{_sfx()}",
                                  "subject_code": "NOT_EXIST"})
        assert r.status_code == 422 and "我方公司" in r.text

    def test_所属公司可留空且显示名回退仓库名(self, box):
        wh = box.warehouse(subject=None)
        assert wh["subject_code"] is None
        assert wh["display_name"] == wh["name"], "未维护所属公司时回退纯仓库名（兼容存量）"

    def test_仓库下拉返回显示名(self, box):
        wh = box.warehouse(subject=SUBJECT2)
        rows = box.client.get("/api/master/options/warehouse", headers=box.admin,
                              params={"keyword": wh["name"]}).json()
        hit = next(r for r in rows if r["id"] == wh["id"])
        assert hit["display_name"] == f"{SUBJECT2}-{wh['name']}"
        assert hit["subject_name"] == "云羲公司"

    def test_修改所属公司(self, box):
        wh = box.warehouse(subject=SUBJECT)
        r = box.client.put(f"/api/master/warehouses/{wh['id']}", headers=box.admin,
                           json={"subject_code": SUBJECT2})
        assert r.status_code == 200, r.text
        assert r.json()["subject_code"] == SUBJECT2 and r.json()["display_name"].startswith(SUBJECT2)


# ===================== N2 单据签约公司 =====================

class TestDocSubject:
    @pytest.mark.parametrize("path", DOC_CREATE_PATHS)
    def test_六类单据缺少签约公司被拒(self, box, path):
        product, wh = box.product(), box.warehouse()
        payload = {"doc_date": "2026-03-01", "items": _items(product)}
        # 各单据的必填差异字段（仓库/供应商/客户）先补齐，确保 422 的原因是"签约公司"
        if path.endswith("orders") and "purchase" in path:
            payload["supplier_id"] = box.supplier()["id"]
        if path.endswith("orders") and "sales" in path:
            payload["customer_id"] = box.customer()["id"]
        if "in-orders" in path or "out-orders" in path:
            payload["warehouse_id"] = wh["id"]
            if "in-orders" in path:
                payload["in_type"] = "其他入库"
            else:
                payload["out_type"] = "其他出库"
        r = box.client.post(path, headers=box.admin, json=payload)
        assert r.status_code == 422, r.text
        assert "签约公司" in r.text

    def test_签约公司取值必须命中字典(self, box):
        product = box.product()
        r = box.client.post("/api/purchase/requests", headers=box.admin,
                            json={"subject_code": "NOPE", "items": _items(product)})
        assert r.status_code == 422 and "我方公司" in r.text

    def test_建单落名称快照且详情可读(self, box):
        doc = box.create_doc("/api/purchase/requests",
                             {"subject_code": SUBJECT, "items": _items(box.product())})
        assert doc["subject_code"] == SUBJECT and doc["subject_name"] == SUBJECT_NAME
        detail = box.client.get(f"/api/purchase/requests/{doc['id']}", headers=box.admin).json()
        assert detail["subject_name"] == SUBJECT_NAME

    def test_修改签约公司(self, box):
        doc = box.create_doc("/api/purchase/requests",
                             {"subject_code": SUBJECT, "items": _items(box.product())})
        r = box.client.put(f"/api/purchase/requests/{doc['id']}", headers=box.admin,
                           json={"subject_code": SUBJECT2})
        assert r.status_code == 200, r.text
        assert r.json()["subject_code"] == SUBJECT2
        assert r.json()["subject_name"] == "云羲公司"
        # 必填字段不允许清空
        r = box.client.put(f"/api/purchase/requests/{doc['id']}", headers=box.admin,
                           json={"subject_code": ""})
        assert r.status_code == 422 and "签约公司" in r.text

    def test_列表可按签约公司筛选(self, box):
        a = box.create_doc("/api/purchase/requests",
                           {"subject_code": SUBJECT, "items": _items(box.product())})
        b = box.create_doc("/api/purchase/requests",
                           {"subject_code": SUBJECT2, "items": _items(box.product())})
        rows = box.client.get("/api/purchase/requests", headers=box.admin,
                              params={"subject_code": SUBJECT2, "page_size": 100}).json()["items"]
        ids = {r["id"] for r in rows}
        assert b["id"] in ids and a["id"] not in ids
        assert all(r["subject_code"] == SUBJECT2 for r in rows)

    def test_导出含签约公司列(self, box):
        doc = box.create_doc("/api/purchase/requests",
                             {"subject_code": SUBJECT, "items": _items(box.product())})
        resp = box.client.get("/api/purchase/requests/export.xlsx", headers=box.admin,
                              params={"keyword": doc["doc_no"]})
        assert resp.status_code == 200, resp.text
        ws = load_workbook(io.BytesIO(resp.content)).active
        rows = list(ws.values)
        header = list(rows[0])
        assert "签约公司" in header
        assert rows[1][header.index("签约公司")] == SUBJECT_NAME
        assert rows[1][header.index("单据编号")] == doc["doc_no"]

    def test_下推继承签约公司(self, box):
        req = box.create_doc("/api/purchase/requests",
                             {"subject_code": SUBJECT2, "items": _items(box.product(), 5)})
        box.client.post(f"/api/purchase/requests/{req['id']}/submit", headers=box.admin)
        box.client.post(f"/api/purchase/requests/{req['id']}/approve", headers=box.admin)
        r = box.client.post(f"/api/purchase/requests/{req['id']}/push", headers=box.admin,
                            json={"supplier_id": box.supplier()["id"]})
        assert r.status_code == 200, r.text
        po = r.json()
        box.docs.setdefault("purchase_orders", []).append(po["id"])
        assert po["subject_code"] == SUBJECT2 and po["subject_name"] == "云羲公司"

    def test_来源单据缺签约公司时下推被拒(self, box, db):
        """历史单据未回填签约公司时，下推必须明确报错，而不是生成空值单据。"""
        from datetime import date

        from app.models_doc import PurchaseRequest, PurchaseRequestItem

        admin_user = db.query(User).filter(User.username == ADMIN_USERNAME).one()
        product = box.product()
        sfx = _sfx()
        pr = PurchaseRequest(doc_no=f"PRX{sfx}", doc_date=date.today(), status="approved",
                             created_by=admin_user.id, subject_code=None)
        db.add(pr)
        db.flush()
        pr.items.append(PurchaseRequestItem(
            seq=1, product_id=product["id"], product_code=product["code"],
            product_name=product["name"], qty=5, unit_price=10, amount=50))
        db.commit()
        box.docs.setdefault("purchase_requests", []).append(pr.id)
        r = box.client.post(f"/api/purchase/requests/{pr.id}/push", headers=box.admin,
                            json={"supplier_id": box.supplier()["id"]})
        assert r.status_code == 422, r.text
        assert "签约公司" in r.text


# ===================== N3 库存明细按公司筛选 =====================

class TestBalanceCompanyFilter:
    def test_库存明细按公司筛选(self, box):
        product = box.product()
        wh_zc, wh_yx = box.warehouse(subject=SUBJECT), box.warehouse(subject=SUBJECT2)
        for wh in (wh_zc, wh_yx):
            doc = box.create_doc("/api/stock/in-orders",
                                 {"subject_code": SUBJECT, "warehouse_id": wh["id"],
                                  "in_type": "其他入库", "items": _items(product, 4)})
            box.client.post(f"/api/stock/in-orders/{doc['id']}/submit", headers=box.admin)
            box.client.post(f"/api/stock/in-orders/{doc['id']}/approve", headers=box.admin)

        rows = box.client.get("/api/stock/balances", headers=box.admin,
                              params={"keyword": product["code"], "subject_code": SUBJECT}).json()["items"]
        assert rows, "按公司筛选应命中 ZC 仓库的结存"
        assert all(r["subject_code"] == SUBJECT for r in rows)
        assert all(r["warehouse_display_name"].startswith(f"{SUBJECT}-") for r in rows)

        rows2 = box.client.get("/api/stock/balances", headers=box.admin,
                               params={"keyword": product["code"], "subject_code": SUBJECT2}).json()["items"]
        assert all(r["subject_code"] == SUBJECT2 for r in rows2)
        assert {r["warehouse_id"] for r in rows2} == {wh_yx["id"]}

    def test_库存结存导出按公司筛选(self, box):
        product = box.product()
        wh = box.warehouse(subject=SUBJECT2)
        doc = box.create_doc("/api/stock/in-orders",
                             {"subject_code": SUBJECT, "warehouse_id": wh["id"],
                              "in_type": "其他入库", "items": _items(product, 3)})
        box.client.post(f"/api/stock/in-orders/{doc['id']}/submit", headers=box.admin)
        box.client.post(f"/api/stock/in-orders/{doc['id']}/approve", headers=box.admin)

        resp = box.client.get("/api/stock/balances/export.xlsx", headers=box.admin,
                              params={"keyword": product["code"], "subject_code": SUBJECT2})
        assert resp.status_code == 200, resp.text
        ws = load_workbook(io.BytesIO(resp.content)).active
        rows = list(ws.values)
        header = list(rows[0])
        assert rows[1][header.index("公司")] == SUBJECT2
        assert rows[1][header.index("仓库")] == f"{SUBJECT2}-{wh['name']}"

        # 换成另一个公司筛选 → 不应再有该行
        resp = box.client.get("/api/stock/balances/export.xlsx", headers=box.admin,
                              params={"keyword": product["code"], "subject_code": SUBJECT})
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert len(list(ws.values)) == 1, "另一公司下不应出现该仓库结存"


# ===================== N4 调拨单调出/调入单位 =====================

class TestTransferUnit:
    def test_单位随仓库所属公司自动带出(self, box):
        wh_from, wh_to = box.warehouse(subject=SUBJECT), box.warehouse(subject=SUBJECT2)
        product = box.product()
        # 先把货入到调出仓，保证调拨过账有库存（审核过账会被库存校验拦截）
        doc = box.create_doc("/api/stock/in-orders",
                             {"subject_code": SUBJECT, "warehouse_id": wh_from["id"],
                              "in_type": "其他入库", "items": _items(product, 5)})
        box.client.post(f"/api/stock/in-orders/{doc['id']}/submit", headers=box.admin)
        box.client.post(f"/api/stock/in-orders/{doc['id']}/approve", headers=box.admin)

        tr = box.create_doc("/api/stock/transfers",
                            {"from_warehouse_id": wh_from["id"], "to_warehouse_id": wh_to["id"],
                             "items": _items(product, 2, 0)})
        assert tr["from_subject_code"] == SUBJECT and tr["from_subject_name"] == SUBJECT_NAME
        assert tr["to_subject_code"] == SUBJECT2 and tr["to_subject_name"] == "云羲公司"

    def test_仓库无所属公司时单位为空(self, box):
        wh1, wh2 = box.warehouse(subject=None), box.warehouse(subject=None)
        tr = box.create_doc("/api/stock/transfers",
                            {"from_warehouse_id": wh1["id"], "to_warehouse_id": wh2["id"],
                             "items": _items(box.product(), 1, 0)})
        assert tr["from_subject_code"] is None and tr["to_subject_code"] is None

    def test_仓库所属公司变更后按新值带出(self, box):
        wh1, wh2 = box.warehouse(subject=SUBJECT), box.warehouse(subject=SUBJECT)
        box.client.put(f"/api/master/warehouses/{wh2['id']}", headers=box.admin,
                       json={"subject_code": SUBJECT2})
        tr = box.create_doc("/api/stock/transfers",
                            {"from_warehouse_id": wh1["id"], "to_warehouse_id": wh2["id"],
                             "items": _items(box.product(), 1, 0)})
        assert tr["from_subject_name"] == SUBJECT_NAME and tr["to_subject_name"] == "云羲公司"


# ===================== 回归：断言辅助 =====================

def test_仓库档案接口字段齐全(box):
    """仓库列表字段必须包含 V2.3 新字段（前端列与筛选依赖）。"""
    box.warehouse(subject=SUBJECT)
    rows = box.client.get("/api/master/warehouses", headers=box.admin).json()["items"]
    assert rows
    for key in ("subject_code", "subject_name", "display_name"):
        assert key in rows[0], f"仓库档案缺少字段 {key}"
