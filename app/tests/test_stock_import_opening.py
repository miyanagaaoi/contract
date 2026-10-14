"""期初库存导入（仅超管）测试。

覆盖 `POST /api/stock/balances/import` 与 `GET /api/stock/balances/import-template.xlsx`：
- 权限：非超管（buyer/keeper）→ 403（权限点方案不满足"仅超管"语义，此处锁死）；
- 期初语义：结存被**设为**导入值、差额写「期初入库」流水；
- 幂等：重复导入同一文件 → 全部跳过；目标 0 与当前 0 → 跳过且不写流水；
- 合并：同物料+仓库多行取最后一行；错误行跳过、有效行提交。
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.database import SessionLocal
from app.init_db import ADMIN_INIT_PASSWORD, ADMIN_USERNAME
from app.main import app
from app.models_auth import Role, User
from app.models_master import Product, ProductType, Uom, Warehouse
from app.models_stock import Stock, StockLedger
from app.security import hash_password

IMPORT_URL = "/api/stock/balances/import"
TEMPLATE_URL = "/api/stock/balances/import-template.xlsx"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def admin_h(client):
    resp = client.post("/api/auth/login",
                       json={"username": ADMIN_USERNAME, "password": ADMIN_INIT_PASSWORD})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['token']}"}


def _user_headers(client, role_code: str) -> dict:
    """建一个非超管账号并登录（角色取种子角色 buyer/keeper）。"""
    username = f"imp_{uuid.uuid4().hex[:8]}"
    with SessionLocal() as db:
        role = db.query(Role).filter(Role.code == role_code).one()
        user = User(username=username, real_name=f"导入测试-{role_code}",
                    password_hash=hash_password("init12345"), status="enabled")
        user.roles.append(role)
        db.add(user)
        db.commit()
    resp = client.post("/api/auth/login",
                       json={"username": username, "password": "init12345"})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['token']}"}


@pytest.fixture()
def env():
    """一组独立物料/仓库，避免与其它用例的结存互相干扰。

    用完**物理清理**：库存明细列表默认按物料编码取第一页（无筛选分页），
    遗留的非零结存行会把 `test_v2_docs` 里"在第一页里找自己那行"的断言挤掉。
    """
    sfx = uuid.uuid4().hex[:6]
    with SessionLocal() as db:
        uom = Uom(code=f"U{sfx}", name=f"个{sfx}", decimals=0)
        db.add(uom)
        db.flush()
        ptype = ProductType(name=f"类型{sfx}")
        db.add(ptype)
        db.flush()
        p1 = Product(code=f"PA{sfx}", name=f"物料A{sfx}", product_type_id=ptype.id, uom_id=uom.id)
        p2 = Product(code=f"PB{sfx}", name=f"物料B{sfx}", product_type_id=ptype.id, uom_id=uom.id)
        wh = Warehouse(code=f"W{sfx}", name=f"仓库{sfx}")
        db.add_all([p1, p2, wh])
        db.commit()
        ids = {"p1": p1.id, "p2": p2.id, "wh": wh.id,
               "p1_code": p1.code, "p2_code": p2.code, "wh_name": wh.name, "sfx": sfx}
    yield ids
    with SessionLocal() as db:
        db.query(StockLedger).filter(StockLedger.product_id.in_([ids["p1"], ids["p2"]]),
                                     StockLedger.doc_type == "stock_opening").delete(
                                     synchronize_session=False)
        db.query(Stock).filter(Stock.product_id.in_([ids["p1"], ids["p2"]])).delete(
            synchronize_session=False)
        db.query(Product).filter(Product.id.in_([ids["p1"], ids["p2"]])).delete(
            synchronize_session=False)
        db.query(Warehouse).filter(Warehouse.id == ids["wh"]).delete(synchronize_session=False)
        db.query(ProductType).filter(ProductType.name == f"类型{sfx}").delete(
            synchronize_session=False)
        db.query(Uom).filter(Uom.code == f"U{sfx}").delete(synchronize_session=False)
        db.commit()


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["物料编码*", "仓库*", "结存数量*", "单价"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, headers: dict, rows: list[list]):
    return client.post(IMPORT_URL, headers=headers,
                       files={"file": ("opening.xlsx", _xlsx(rows),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _stock_qty(product_id: int, warehouse_id: int) -> Decimal:
    with SessionLocal() as db:
        stock = (db.query(Stock).filter(Stock.product_id == product_id,
                                        Stock.warehouse_id == warehouse_id).first())
        return Decimal(stock.qty or 0) if stock else Decimal("0")


def _opening_ledgers(product_id: int, warehouse_id: int) -> list[StockLedger]:
    with SessionLocal() as db:
        return (db.query(StockLedger)
                .filter(StockLedger.product_id == product_id,
                        StockLedger.warehouse_id == warehouse_id,
                        StockLedger.biz_type == "期初入库")
                .order_by(StockLedger.id).all())


class TestPermission:
    def test_非超管_导入与模板都_403(self, client):
        for role in ("buyer", "keeper"):
            headers = _user_headers(client, role)
            assert client.post(IMPORT_URL, headers=headers,
                               files={"file": ("a.xlsx", _xlsx([]), "application/octet-stream")}
                               ).status_code == 403
            assert client.get(TEMPLATE_URL, headers=headers).status_code == 403

    def test_未登录_401(self, client):
        assert client.get(TEMPLATE_URL).status_code == 401


class TestImport:
    def test_两有效一错误_生效与错误报告(self, client, admin_h, env):
        rows = [[env["p1_code"], env["wh_name"], 10, 5],
                [env["p2_code"], env["wh_name"], 3, None],
                ["NOPE-404", env["wh_name"], 7, None]]
        resp = _post(client, admin_h, rows)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["total"] == 3
        assert data["imported"] == 2
        assert data["skipped"] == 0
        assert len(data["errors"]) == 1
        assert data["errors"][0]["row"] == 4
        assert "不存在" in data["errors"][0]["reason"]

        assert _stock_qty(env["p1"], env["wh"]) == Decimal("10")
        assert _stock_qty(env["p2"], env["wh"]) == Decimal("3")
        ledgers = _opening_ledgers(env["p1"], env["wh"])
        assert len(ledgers) == 1
        assert Decimal(ledgers[0].qty_change) == Decimal("10")
        assert Decimal(ledgers[0].qty_after) == Decimal("10")
        assert Decimal(ledgers[0].unit_price) == Decimal("5")

    def test_重复导入_幂等全跳过(self, client, admin_h, env):
        rows = [[env["p1_code"], env["wh_name"], 12, 5]]
        assert _post(client, admin_h, rows).json()["imported"] == 1
        data = _post(client, admin_h, rows).json()          # 同一文件再导一次
        assert data["imported"] == 0
        assert data["skipped"] == 1
        assert len(_opening_ledgers(env["p1"], env["wh"])) == 1   # 不新增流水

    def test_导入数量0_当前也为0_跳过不写流水(self, client, admin_h, env):
        data = _post(client, admin_h, [[env["p2_code"], env["wh_name"], 0, None]]).json()
        assert data["imported"] == 0
        assert data["skipped"] == 1
        assert _opening_ledgers(env["p2"], env["wh"]) == []

    def test_同物料同仓库多行_取最后一行并提示合并(self, client, admin_h, env):
        rows = [[env["p1_code"], env["wh_name"], 8, None],
                [env["p1_code"], env["wh_name"], 20, 2]]
        data = _post(client, admin_h, rows).json()
        assert data["merged"] == 1
        assert data["imported"] == 1
        assert _stock_qty(env["p1"], env["wh"]) == Decimal("20")

    def test_仓库支持主体码前缀显示名(self, client, admin_h, env):
        data = _post(client, admin_h,
                     [[env["p2_code"], f"ZC-{env['wh_name']}", 6, None]]).json()
        assert data["imported"] == 1
        assert _stock_qty(env["p2"], env["wh"]) == Decimal("6")

    def test_结存调减_负差额流水(self, client, admin_h, env):
        _post(client, admin_h, [[env["p2_code"], env["wh_name"], 9, None]])
        _post(client, admin_h, [[env["p2_code"], env["wh_name"], 4, None]])   # 目标 4 → 差额 -5
        ledgers = _opening_ledgers(env["p2"], env["wh"])
        assert len(ledgers) == 2
        assert Decimal(ledgers[1].qty_change) == Decimal("-5")
        assert Decimal(ledgers[1].qty_after) == Decimal("4")
        assert _stock_qty(env["p2"], env["wh"]) == Decimal("4")

    def test_非xlsx_422(self, client, admin_h):
        resp = client.post(IMPORT_URL, headers=admin_h,
                           files={"file": ("a.xls", b"not xlsx", "application/octet-stream")})
        assert resp.status_code == 422

    def test_模板下载_200且xlsx(self, client, admin_h):
        resp = client.get(TEMPLATE_URL, headers=admin_h)
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        # 能被 openpyxl 解析即证明是合法 xlsx
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert [c.value for c in ws[1]][:3] == ["物料编码*", "仓库*", "结存数量*"]
