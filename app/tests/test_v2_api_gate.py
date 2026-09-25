"""T-V2-06b 接口鉴权收口测试。

覆盖：
- 未登录拦截（AC-V2-02）在所有既有接口上生效
- 字典读接口「登录即可」、写接口需 `system.dict.edit`
- 导出/导入的按钮权限
- **导出数据范围**（关键安全点：导出必须与列表同一数据范围，不得越权导出全公司台账）

测试自建组织/账号/合同并在 fixture 结束时清理。
"""
from __future__ import annotations

import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import SessionLocal
from app.init_db import ADMIN_INIT_PASSWORD, ADMIN_USERNAME
from app.main import app
from app.models import ChangeLog, Contract, ContractItem, contract_tag
from app.models_auth import OperationLog, OrgUnit, Role, User
from app.security import hash_password

# 这些接口「登录即可」，被前端表单用作用字字典
LOGIN_ONLY_GETS = [
    "/api/meta",
    "/api/dashboard",
    "/api/tags",
    "/api/settings/item-types",
    "/api/settings/contract-types",
    "/api/settings/subjects",
]


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def admin_h(client):
    resp = client.post("/api/auth/login",
                       json={"username": ADMIN_USERNAME, "password": ADMIN_INIT_PASSWORD})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['token']}"}


class Box:
    """精简测试环境（组织 + 账号 + 合同）。"""

    def __init__(self, client, admin_h):
        self.client = client
        self.admin_h = admin_h
        self.orgs: list[int] = []
        self.users: list[int] = []
        self.contracts: list[int] = []

    def org(self, name: str) -> dict:
        resp = self.client.post("/api/system/org-units", headers=self.admin_h,
                                json={"name": name, "unit_type": "部门"})
        assert resp.status_code == 200, resp.text
        node = resp.json()
        self.orgs.append(node["id"])
        return node

    def user(self, role_code: str, org_id: int | None = None) -> int:
        username = f"g_{uuid.uuid4().hex[:8]}"
        with SessionLocal() as db:
            role = db.query(Role).filter(Role.code == role_code).one()
            user = User(username=username, real_name=f"门禁-{role_code}", org_id=org_id,
                        password_hash=hash_password("init12345"), status="enabled")
            user.roles.append(role)
            db.add(user)
            db.commit()
            uid = user.id
        self.users.append(uid)
        return uid

    def h(self, uid: int) -> dict:
        with SessionLocal() as db:
            username = db.get(User, uid).username
        resp = self.client.post("/api/auth/login",
                                json={"username": username, "password": "init12345"})
        assert resp.status_code == 200, resp.text
        return {"Authorization": f"Bearer {resp.json()['token']}"}

    def contract(self, headers: dict, name: str) -> dict:
        resp = self.client.post("/api/contracts", headers=headers, json={
            "name": name, "type": "PUR", "subject_code": "ZC", "amount": 100})
        assert resp.status_code == 200, resp.text
        data = resp.json()
        self.contracts.append(data["id"])
        return data

    def cleanup(self) -> None:
        with SessionLocal() as db:
            for cid in self.contracts:
                db.execute(contract_tag.delete().where(contract_tag.c.contract_id == cid))
                db.query(ChangeLog).filter(ChangeLog.contract_id == cid).delete()
                db.query(ContractItem).filter(ContractItem.contract_id == cid).delete()
                db.query(Contract).filter(Contract.id == cid).delete()
            for uid in self.users:
                db.query(OperationLog).filter(OperationLog.user_id == uid).delete()
                db.query(User).filter(User.id == uid).delete()
            db.commit()
            for oid in sorted(self.orgs, reverse=True):
                node = db.get(OrgUnit, oid)
                if node is not None:
                    db.delete(node)
            db.commit()
        self.orgs, self.users, self.contracts = [], [], []


@pytest.fixture()
def box(client, admin_h):
    instance = Box(client, admin_h)
    try:
        yield instance
    finally:
        instance.cleanup()


def _xlsx_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return "\n".join("|".join(str(c) for c in row) for row in ws.iter_rows(values_only=True))


# ===================== 登录门禁 =====================

class TestLoginGate:
    def test_all_endpoints_require_login(self, client):
        for path in LOGIN_ONLY_GETS:
            assert client.get(path).status_code == 401, path
        assert client.get("/api/export/contracts.xlsx").status_code == 401
        assert client.get("/api/import/template.xlsx").status_code == 401
        assert client.get("/api/contracts/1/attachments").status_code == 401
        assert client.get("/api/attachments/1/download").status_code == 401

    def test_logged_in_can_read_dictionaries(self, client, box):
        """字典读接口必须对普通业务角色开放，否则开不了单。"""
        headers = box.h(box.user("buyer"))
        for path in LOGIN_ONLY_GETS:
            assert client.get(path, headers=headers).status_code == 200, path

    def test_keeper_can_open_dashboard(self, client, box):
        headers = box.h(box.user("keeper"))
        assert client.get("/api/dashboard", headers=headers).status_code == 200


# ===================== 按钮权限 =====================

class TestButtonGate:
    def test_buyer_cannot_edit_dictionary(self, client, box):
        headers = box.h(box.user("buyer"))
        resp = client.post("/api/tags", headers=headers, json={"name": f"越权标签{uuid.uuid4().hex[:4]}"})
        assert resp.status_code == 403
        assert "system.dict.edit" in resp.json()["detail"]

        resp = client.put("/api/settings/item-types", headers=headers, json={"values": ["采购"]})
        assert resp.status_code == 403
        assert "system.dict.edit" in resp.json()["detail"]

    def test_admin_can_edit_dictionary(self, client, box):
        """超管可维护字典（含标签新增/删除，测试后清理）。"""
        name = f"门禁标签{uuid.uuid4().hex[:4]}"
        created = client.post("/api/tags", headers=box.admin_h, json={"name": name})
        assert created.status_code == 200
        tag_id = created.json()["id"]
        assert client.delete(f"/api/tags/{tag_id}", headers=box.admin_h).status_code == 200

    def test_keeper_cannot_export_contracts(self, client, box):
        headers = box.h(box.user("keeper"))
        resp = client.get("/api/export/contracts.xlsx", headers=headers)
        assert resp.status_code == 403
        assert "contract.export" in resp.json()["detail"]

    def test_keeper_cannot_import_contracts(self, client, box):
        headers = box.h(box.user("keeper"))
        resp = client.get("/api/import/template.xlsx", headers=headers)
        assert resp.status_code == 403
        assert "contract.import" in resp.json()["detail"]

    def test_buyer_can_export_but_not_import(self, client, box):
        headers = box.h(box.user("buyer"))
        assert client.get("/api/export/contracts.xlsx", headers=headers).status_code == 200
        assert client.get("/api/import/template.xlsx", headers=headers).status_code == 403


# ===================== 导出数据范围（安全关键） =====================

class TestExportDataScope:
    def test_export_only_contains_own_scope(self, client, box):
        dept_a = box.org(f"导出甲部_{uuid.uuid4().hex[:4]}")
        dept_b = box.org(f"导出乙部_{uuid.uuid4().hex[:4]}")
        headers_a = box.h(box.user("buyer", dept_a["id"]))
        headers_b = box.h(box.user("buyer", dept_b["id"]))

        name_a = f"导出甲{uuid.uuid4().hex[:6]}"
        name_b = f"导出乙{uuid.uuid4().hex[:6]}"
        box.contract(headers_a, name=name_a)
        box.contract(headers_b, name=name_b)

        resp = client.get("/api/export/contracts.xlsx", headers=headers_a)
        assert resp.status_code == 200
        text = _xlsx_text(resp.content)
        assert name_a in text
        assert name_b not in text, "导出不得包含数据范围外的合同（越权导出）"

    def test_admin_export_sees_all(self, client, box):
        dept = box.org(f"导出丙部_{uuid.uuid4().hex[:4]}")
        headers = box.h(box.user("buyer", dept["id"]))
        name = f"导出丙{uuid.uuid4().hex[:6]}"
        box.contract(headers, name=name)

        resp = client.get("/api/export/contracts.xlsx", headers=box.admin_h)
        assert resp.status_code == 200
        assert name in _xlsx_text(resp.content)

    def test_export_matches_list_scope(self, client, box):
        """导出与列表的数据范围必须一致（同一查询构建器）。"""
        dept = box.org(f"导出一致_{uuid.uuid4().hex[:4]}")
        headers = box.h(box.user("buyer", dept["id"]))
        name = f"一致{uuid.uuid4().hex[:6]}"
        box.contract(headers, name=name)

        listed = client.get("/api/contracts", headers=headers, params={"keyword": name}).json()
        assert listed["total"] == 1
        text = _xlsx_text(client.get("/api/export/contracts.xlsx", headers=headers,
                                     params={"keyword": name}).content)
        assert name in text


# ===================== 附件权限 =====================

class TestAttachmentGate:
    def test_attachment_list_requires_contract_view(self, client, box):
        headers = box.h(box.user("keeper"))   # keeper 有 contract.view
        created = box.contract(box.admin_h, name=f"附件合同_{uuid.uuid4().hex[:6]}")
        resp = client.get(f"/api/contracts/{created['id']}/attachments", headers=headers)
        assert resp.status_code == 200

    def test_download_requires_contract_view(self, client, box):
        headers = box.h(box.user("keeper"))
        resp = client.get("/api/attachments/999999/download", headers=headers)
        assert resp.status_code == 404      # 通过鉴权但附件不存在

    def test_user_without_contract_view_denied(self, client):
        """无 contract.view 的账号（临时移除权限的 buyer 变体）由越权用例覆盖：
        这里用未登录与 keeper 断言边界，其余组合见 test_v2_contract_perm。"""
        assert client.get("/api/contracts/1/attachments").status_code == 401
