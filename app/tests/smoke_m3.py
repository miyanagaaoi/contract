"""CTMS V2.0 · M3 验收冒烟（销售线 + 盘点 + 导出 + 打印）。

覆盖：销售申请→销售订单→出库过账闭环、出库负库存拦截、盘点全盘与盘盈/盘亏自动调整单、
单据与库存导出（条数=当前筛选）、单据 A4 打印页。

用法（先启动后端）：
    app\\.venv\\Scripts\\python.exe app\\tests\\smoke_m3.py

可选环境变量：`CTMS_SMOKE_BASE`（默认 http://127.0.0.1:8010/api）
"""
from __future__ import annotations

import io
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

BASE = os.environ.get("CTMS_SMOKE_BASE", "http://127.0.0.1:8010/api")
ADMIN_USER, ADMIN_PWD = "admin", "admin12345"

PASSED: list[str] = []
FAILED: list[str] = []
TS = int(time.time()) % 100000


def ok(name: str) -> None:
    PASSED.append(name)
    print(f"  [PASS] {name}")


def req(method: str, path: str, body=None, params=None, expect=(200,), token=None, binary=False):
    url = BASE + path
    if params:
        url += "?" + urllib.parse.urlencode({k: v for k, v in params.items() if v is not None})
    data = json.dumps(body, ensure_ascii=True).encode("utf-8") if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(request) as resp:
            raw = resp.read()
            if binary:
                return resp.status, raw
            try:
                return resp.status, json.loads(raw.decode("utf-8") or "null")
            except (json.JSONDecodeError, UnicodeDecodeError):
                return resp.status, raw
    except urllib.error.HTTPError as exc:
        raw = exc.read()
        if binary:
            return exc.code, raw
        try:
            detail = json.loads(raw.decode("utf-8") or "null")
        except (json.JSONDecodeError, UnicodeDecodeError):
            detail = {"raw": raw[:200]}
        if exc.code not in expect:
            raise AssertionError(f"{method} {path} -> {exc.code} {detail}")
        return exc.code, detail


def login(username: str, password: str) -> str:
    code, data = req("POST", "/auth/login", {"username": username, "password": password})
    assert code == 200 and data.get("token"), f"登录失败：{username}"
    return data["token"]


class Cleaner:
    def __init__(self) -> None:
        self.products: list[int] = []
        self.types: list[int] = []
        self.uoms: list[int] = []
        self.warehouses: list[int] = []
        self.customers: list[int] = []
        self.contracts: list[int] = []
        self.docs: dict[str, list[int]] = {}
        self.usernames: list[str] = []

    def cleanup(self, admin: str) -> None:
        for cid in self.contracts:
            try:
                req("DELETE", f"/contracts/{cid}", params={"reason": "M3 冒烟归档"}, token=admin)
            except Exception:  # noqa: BLE001
                pass
        for path, ids in (("/master/products", self.products), ("/master/uoms", self.uoms),
                          ("/master/product-types", list(reversed(self.types))),
                          ("/master/warehouses", self.warehouses),
                          ("/master/customers", self.customers)):
            for obj_id in ids:
                try:
                    req("DELETE", f"{path}/{obj_id}", token=admin)
                except Exception:  # noqa: BLE001
                    pass
        try:
            from sqlalchemy import text

            from app import models_auth, models_doc, models_master, models_stock  # noqa: F401
            from app.database import SessionLocal
            from app.models import ChangeLog
            from app.models_auth import OperationLog, User
            from app.models_doc import DOC_MODELS
            from app.models_stock import Stock, StockLedger

            with SessionLocal() as db:
                for table, ids in self.docs.items():
                    for doc_id in ids:
                        db.execute(text(f"DELETE FROM {table} WHERE id = :i"), {"i": doc_id})
                db.query(ChangeLog).filter(ChangeLog.object_type.in_(list(DOC_MODELS))).delete(
                    synchronize_session=False)
                for pid in self.products:
                    db.query(StockLedger).filter(StockLedger.product_id == pid).delete()
                    db.query(Stock).filter(Stock.product_id == pid).delete()
                for username in self.usernames:
                    row = db.query(User).filter(User.username == username).first()
                    if row is not None:
                        db.query(OperationLog).filter(OperationLog.user_id == row.id).delete()
                        db.delete(row)
                db.commit()
        except Exception as exc:  # noqa: BLE001
            print(f"  [WARN] 清理残留数据失败：{exc}")


def mk_user(admin: str, cleaner: Cleaner, role_code: str, real_name: str) -> str:
    username = f"m3_{role_code[:6]}_{TS}_{len(cleaner.usernames)}"
    _, roles = req("GET", "/system/roles", token=admin)
    role_id = next(r["id"] for r in roles["items"] if r["code"] == role_code)
    req("POST", "/system/users", {"username": username, "real_name": real_name,
                                  "password": "init12345", "role_ids": [role_id]}, token=admin)
    cleaner.usernames.append(username)
    return login(username, "init12345")


def track(cleaner: Cleaner, kind: str, doc: dict) -> dict:
    cleaner.docs.setdefault(kind, []).append(doc["id"])
    return doc


def xlsx_rows(blob: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(blob), read_only=True)
    return list(wb.active.iter_rows(values_only=True))


def main() -> int:
    print(f"M3 验收冒烟 → {BASE}")
    admin = login(ADMIN_USER, ADMIN_PWD)
    c = Cleaner()
    try:
        run(admin, c)
    except AssertionError as exc:
        FAILED.append(str(exc))
        print(f"  [FAIL] {exc}")
    finally:
        c.cleanup(admin)
    print(f"\nM3 冒烟结果：{len(PASSED)} 项通过，{len(FAILED)} 项失败")
    for item in FAILED:
        print(f"  - {item}")
    return 1 if FAILED else 0


def run(admin: str, c: Cleaner) -> None:
    seller = mk_user(admin, c, "seller", "M3销售员")
    manager = mk_user(admin, c, "sales_manager", "M3销售主管")
    keeper = mk_user(admin, c, "keeper", "M3仓管员")
    buyer = mk_user(admin, c, "buyer", "M3采购员")   # 无销售单据权限，用于打印越权校验

    # ---------- 主数据 ----------
    _, uom = req("POST", "/master/uoms", {"code": f"MU{TS % 1000:03d}", "name": f"件{TS}",
                                          "decimals": 2}, token=admin)
    c.uoms.append(uom["id"])
    _, ptype = req("POST", "/master/product-types", {"name": f"M3类型{TS}", "code": f"Z{TS % 100:02d}"},
                   token=admin)
    c.types.append(ptype["id"])
    _, product = req("POST", "/master/products", {"name": f"M3物料{TS}", "product_type_id": ptype["id"],
                                                  "uom_id": uom["id"], "default_price": 20,
                                                  "safety_stock": 50}, token=admin)
    c.products.append(product["id"])
    _, warehouse = req("POST", "/master/warehouses", {"code": f"MW{TS % 1000:03d}",
                                                      "name": f"M3仓{TS}"}, token=admin)
    c.warehouses.append(warehouse["id"])
    _, customer = req("POST", "/master/customers", {"name": f"M3客户{TS}"}, token=admin)
    c.customers.append(customer["id"])
    _, contract = req("POST", "/contracts", {"name": f"M3销售合同{TS}", "type": "SAL",
                                             "subject_code": "ZC", "amount": 8000}, token=admin)
    c.contracts.append(contract["id"])
    items = [{"product_id": product["id"], "qty": 30, "unit_price": 20}]

    # ---------- 备货：入库 30 ----------
    _, stock_in = req("POST", "/stock/in-orders", {
        "warehouse_id": warehouse["id"], "in_type": "其他入库", "items": items}, token=keeper)
    track(c, "stock_in_orders", stock_in)
    req("POST", f"/stock/in-orders/{stock_in['id']}/submit", token=keeper)
    req("POST", f"/stock/in-orders/{stock_in['id']}/approve", token=admin)

    # ---------- 销售闭环：申请 → 订单 → 出库 ----------
    _, sreq = req("POST", "/sales/requests", {"customer_id": customer["id"], "items": items,
                                              "contract_id": contract["id"]}, token=seller)
    track(c, "sales_requests", sreq)
    assert sreq["customer_name_text"] == customer["name"] and sreq["total_amount"] == 600.0
    req("POST", f"/sales/requests/{sreq['id']}/submit", token=seller)
    code, _ = req("POST", f"/sales/requests/{sreq['id']}/approve", expect=(403,), token=seller)
    assert code == 403, "销售员不应有审核权限"
    req("POST", f"/sales/requests/{sreq['id']}/approve", token=manager)

    _, so = req("POST", f"/sales/requests/{sreq['id']}/push",
                {"customer_id": customer["id"], "items": [{"qty": 20}]}, token=seller)
    track(c, "sales_orders", so)
    assert so["status"] == "draft" and so["items"][0]["qty"] == 20
    detail = req("GET", f"/sales/requests/{sreq['id']}", token=seller)[1]
    assert detail["items"][0]["ordered_qty"] == 20
    _, rest = req("POST", f"/sales/requests/{sreq['id']}/push",
                  {"customer_id": customer["id"]}, token=seller)
    track(c, "sales_orders", rest)
    assert rest["items"][0]["qty"] == 10, "默认下推剩余 10"
    ok("AC-V2-16 同构 销售申请下推销售订单（20 + 剩余 10）")

    req("POST", f"/sales/orders/{so['id']}/submit", token=seller)
    req("POST", f"/sales/orders/{so['id']}/approve", token=manager)
    code, out = req("POST", f"/sales/orders/{so['id']}/push",
                    {"warehouse_id": warehouse["id"]}, token=manager)
    assert code == 200, out
    track(c, "stock_out_orders", out)
    req("POST", f"/stock/out-orders/{out['id']}/submit", token=keeper)
    _, posted = req("POST", f"/stock/out-orders/{out['id']}/approve", token=admin)
    assert posted["approved_at"] and posted["posted"] is True

    ledger = req("GET", "/stock/ledger", params={"product_id": product["id"],
                                                 "warehouse_id": warehouse["id"]}, token=keeper)[1]
    assert ledger["balance"] == 10.0, ledger["balance"]
    so_detail = req("GET", f"/sales/orders/{so['id']}", token=seller)[1]
    assert so_detail["items"][0]["shipped_qty"] == 20, "回写出库数量"
    ok("AC-V2-20 同构 销售出库过账（结存 30→10）并回写已出库数量")

    # ---------- 负库存拦截 ----------
    _, big = req("POST", "/stock/out-orders", {
        "warehouse_id": warehouse["id"], "out_type": "其他出库",
        "items": [{"product_id": product["id"], "qty": 99, "unit_price": 20}]}, token=keeper)
    track(c, "stock_out_orders", big)
    req("POST", f"/stock/out-orders/{big['id']}/submit", token=keeper)
    code, body = req("POST", f"/stock/out-orders/{big['id']}/approve", expect=(422,), token=admin)
    assert "库存不足" in body["detail"]
    balance = req("GET", "/stock/ledger", params={"product_id": product["id"],
                                                  "warehouse_id": warehouse["id"]}, token=keeper)[1]["balance"]
    assert balance == 10.0
    ok("AC-V2-21 出库负库存拦截（库存与状态不变）")

    # ---------- 盘点：盘亏 → 自动生成盘亏出库单 ----------
    _, take = req("POST", "/stock/takes", {"warehouse_id": warehouse["id"], "take_type": "full"},
                  token=keeper)
    track(c, "stock_takes", take)
    _, generated = req("POST", f"/stock/takes/{take['id']}/generate", {}, token=keeper)
    row = next(i for i in generated["items"] if i["product_id"] == product["id"])
    assert row["book_qty"] == 10.0
    req("PUT", f"/stock/takes/{take['id']}/count",
        {"counts": [{"id": row["id"], "actual_qty": 7, "diff_reason": "破损 3 件"}]}, token=keeper)
    req("POST", f"/stock/takes/{take['id']}/submit", token=keeper)
    _, done = req("POST", f"/stock/takes/{take['id']}/approve", token=admin)
    assert done["generated_out_no"], done
    track(c, "stock_out_orders", {"id": done["generated_out_id"]})
    balance = req("GET", "/stock/ledger", params={"product_id": product["id"],
                                                  "warehouse_id": warehouse["id"]}, token=keeper)[1]["balance"]
    assert balance == 7.0, balance
    ok(f"AC-V2-27/28 盘点全盘与盘亏自动调整（{done['generated_out_no']}，结存 10→7）")

    # ---------- 盘盈 ----------
    _, take2 = req("POST", "/stock/takes", {"warehouse_id": warehouse["id"], "take_type": "full"},
                   token=keeper)
    track(c, "stock_takes", take2)
    _, gen2 = req("POST", f"/stock/takes/{take2['id']}/generate", {}, token=keeper)
    row2 = next(i for i in gen2["items"] if i["product_id"] == product["id"])
    req("PUT", f"/stock/takes/{take2['id']}/count",
        {"counts": [{"id": row2["id"], "actual_qty": 9}]}, token=keeper)
    req("POST", f"/stock/takes/{take2['id']}/submit", token=keeper)
    _, done2 = req("POST", f"/stock/takes/{take2['id']}/approve", token=admin)
    assert done2["generated_in_no"], done2
    track(c, "stock_in_orders", {"id": done2["generated_in_id"]})
    balance = req("GET", "/stock/ledger", params={"product_id": product["id"],
                                                  "warehouse_id": warehouse["id"]}, token=keeper)[1]["balance"]
    assert balance == 9.0, balance
    ok(f"AC-V2-29 盘盈自动调整（{done2['generated_in_no']}，结存 7→9）")

    # ---------- 安全库存预警 ----------
    alerts = req("GET", "/stock/balances", params={"keyword": product["code"],
                                                   "below_safety": True}, token=keeper)[1]
    assert alerts["total"] >= 1 and alerts["items"][0]["below_safety"] is True
    ok("T-V2-34 安全库存预警（below_safety 筛选可用）")

    # ---------- 一致性：结存 = 流水累计 ----------
    recalc = req("POST", "/stock/recalc", token=keeper)[1]
    assert recalc["consistent"] is True
    ok("AC-V2-26 结存与流水一致（recalc）")

    # ---------- 导出（AC-V2-40） ----------
    status, blob = req("GET", "/sales/orders/export.xlsx", params={"keyword": so["doc_no"]},
                       token=seller, binary=True)
    assert status == 200 and blob[:2] == b"PK"
    rows = xlsx_rows(blob)
    assert rows[0][0] == "单据编号" and len(rows) - 1 == 1 and rows[1][0] == so["doc_no"]
    ok("AC-V2-40 销售订单导出（条数=当前筛选）")

    status, blob = req("GET", "/stock/balances/export.xlsx", params={"keyword": product["code"]},
                       token=keeper, binary=True)
    assert status == 200 and blob[:2] == b"PK"
    rows = xlsx_rows(blob)
    assert rows[0][0] == "物料编码" and rows[1][6] == 9.0
    ok("AC-V2-40 库存结存导出")

    # ---------- 打印（T-V2-37） ----------
    status, html = req("GET", f"/stock/out-orders/{out['id']}/print", token=keeper, binary=True)
    text = html.decode("utf-8", "replace")
    assert status == 200, f"打印接口状态 {status}"
    assert "<!DOCTYPE html>" in text, "打印页不是 HTML"
    assert out["doc_no"] in text, f"打印页缺少单号 {out['doc_no']}"
    assert "window.print()" in text, "打印页缺少打印按钮"
    assert "领料" in text, "出库单打印页缺少领料签字栏"
    status, _ = req("GET", f"/sales/orders/{so['id']}/print", expect=(403,), token=buyer)
    assert status == 403, "采购员无销售订单查看权限，打印也应被拦"
    ok("T-V2-37 单据 A4 打印（按权限鉴权）")

    # ---------- 合同关联单据汇总（含销售单据） ----------
    related = req("GET", f"/contracts/{contract['id']}/related-docs", token=admin)[1]
    assert related["summary"]["sales_order_count"] >= 1
    assert related["summary"]["sales_order_amount"] == 600.0
    ok("AC-V2-32 同构 合同关联单据汇总（销售订单 600 元）")


if __name__ == "__main__":
    sys.exit(main())
