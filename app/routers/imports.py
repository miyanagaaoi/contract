"""合同批量导入 API（MVP2 需求② + MVP3 对齐）。

- GET  /api/import/template.xlsx            下载系统导入模板
- POST /api/import/contracts                上传模板，仅新建；错误行整条跳过并返回报告

MVP3 对齐要点：
- 新增"主体码"列（我方公司 ZC/YX，可空=默认第一个可用主体）
- "合同编号"可留空 → 按 类型码+主体码+年份+月份+6位序号 自动生成（同一文件按主档顺序分配）
- 行项明细以"主档序号"关联主档（兼容旧列"合同编号(对应主档)"手填编号方式）
- 类型列支持：新类型名（采购/支出…）、类型代码（PUR…）、旧名（采购/销售）自动映射
"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ..database import get_db
from ..dicts import get_enabled_contract_types, get_subjects, type_code_of, type_label_of
from ..models import STATUSES, Contract
from ..numbering import next_number
from .contracts import create_contract

router = APIRouter(prefix="/api/import", tags=["import"])

MAIN_HEAD = ["序号", "合同编号", "合同名称", "类型", "主体码", "甲方", "乙方", "签订日期",
             "金额", "已付金额", "经办人", "状态", "所属框架编号", "是否框架",
             "标签(逗号分隔)", "备注", "质保金额", "质保比例%", "质保生效日期", "质保期限(月)"]
ITEMS_HEAD = ["主档序号", "序号", "类型", "名称", "规格型号", "数量", "单价", "备注"]

MAIN_SAMPLE = [
    [1, None, "示例：新购设备合同（自动编号）", "采购/支出", "ZC", "本公司（甲方示例）", "XX 供应商", "2026-09-01",
     None, 0, "张三", "内部审批中", None, "否", "采购,项目A", "", 6600, 5, "2026-09-01", 12],
    [2, None, "示例：产品销售（自动编号）", "SAL", "ZC", "客户 M 公司", "本公司（乙方示例）", "2026-09-05",
     None, 6000, "李四", "到货", None, "否", "销售", "", 300, 5, "2026-09-05", 12],
]
ITEMS_SAMPLE = [
    [1, 1, "采购", "X 系列设备", "X-2000", 10, 12000, "含安装调试"],
    [1, 2, "采购", "配套耗材", "HC-02", 2, 6000, ""],
    [2, 1, "销售", "X 设备(整机)", "X-2000", 1, 6000, "已收款"],
]

LIMIT_MAIN = 2000


def _style(ws, headers):
    ws.append(headers)
    from openpyxl.styles import Font, PatternFill

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "A2"


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "合同(主档)"
    _style(ws, MAIN_HEAD)
    for row in MAIN_SAMPLE:
        ws.append(row)
    ws2 = wb.create_sheet("行项明细")
    _style(ws2, ITEMS_HEAD)
    for row in ITEMS_SAMPLE:
        ws2.append(row)
    note = wb.create_sheet("填写说明")
    tips = [
        "填写说明：",
        "1. 工作表『合同(主档)』每行一个合同：『序号』填 1,2,3…（行项明细用它在明细表引用该合同）；示例行删除后填写。",
        "2. 『合同编号』：留空 = 系统自动生成（类型码+主体码+年份+月份+6位序号，年度递增，如 PURZC202609000001）；也可手填（需唯一、不重复）。",
        "3. 『类型』必填：填名称（销售/收入、采购/支出、合作/战略协议、劳动/人事、金融/投融资、保密协议）或代码（SAL/PUR/COO/LAB/FIN/NDA）。",
        "4. 『主体码』：我方公司，ZC=智澈公司 / YX=云羲公司（可空，默认第一个可用主体）。",
        "5. 『行项明细』工作表：『主档序号』填其归属合同在“合同(主档)”里的序号；同一合同的行项序号连续；无行项的合同，金额必须手填。",
        "6. 金额留空=按行项合计自动；金额/数量/单价/比例只填数字；日期填 YYYY-MM-DD。",
        "7. 标签列多个用英文逗号分隔，不存在会自动创建；类型会自动带【类型名】标签。",
        "8. 状态可选：内部审批中/集团审批中/已签订/付款中/发货/到货/已终止（留空默认内部审批中）。",
        "9. 是否框架填是/否；所属框架编号填已存在框架合同编号（子合同自动挂到其下）。",
        "10. 同一合同编号只出现一次：文件内重复、与库内重复、必填缺失、格式错误的行会整条跳过并写入错误报告。",
        "11. 上线前请先用真实数据试 3~5 条。",
    ]
    for t in tips:
        note.append([t])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _clean(s) -> str:
    return "" if s is None else str(s).strip()


def _to_date(v):
    if v is None or _clean(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _clean(v)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"日期格式应为 YYYY-MM-DD：{s}")


def _to_dec(v, field: str, allow_none: bool = False) -> Decimal | None:
    if v is None or _clean(v) == "":
        if allow_none:
            return None
        raise ValueError(f"字段 {field} 必填")
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"字段 {field} 应为数字：{v}")


def _read_rows(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_clean(h) for h in rows[0]]
    data = []
    for idx, row in enumerate(rows[1:], start=2):
        if all((v is None or _clean(v) == "") for v in row):
            continue
        data.append((idx, {headers[j]: row[j] for j in range(min(len(headers), len(row)))}))
    return headers, data


def _ordinal(row: dict, fallback: int) -> int:
    v = _clean(row.get("序号"))
    if v.isdigit() and int(v) > 0:
        return int(v)
    return fallback


def _resolve_type(db: Session, raw: str) -> str:
    """类型 → 标准代码；支持 名称/代码/旧名。无匹配返回 None（由调用方报错）。"""
    if not raw:
        return None
    return type_code_of(db, raw)


def _parse_main(db: Session, row: dict) -> dict:
    name = _clean(row.get("合同名称"))
    if not name:
        raise ValueError("合同名称必填")
    raw_type = _clean(row.get("类型"))
    if not raw_type:
        raise ValueError("类型必填（名称或代码，见填写说明第 3 条）")
    code = type_code_of(db, raw_type)
    if code is None:
        raise ValueError(f"无法识别的类型：{raw_type}（支持 SAL/PUR/COO/LAB/FIN/NDA）")
    status = _clean(row.get("状态")) or "内部审批中"
    amount = _to_dec(row.get("金额"), "金额", allow_none=True)
    payload: dict = {
        "contract_no": _clean(row.get("合同编号")),  # 可为空→自动编号
        "name": name,
        "type": type_label_of(db, code),
        "party_a": _clean(row.get("甲方")),
        "party_b": _clean(row.get("乙方")),
        "sign_date": _to_date(row.get("签订日期")),
        "paid_amount": _to_dec(row.get("已付金额"), "已付金额", allow_none=True) or Decimal("0"),
        "owner_name": _clean(row.get("经办人")) or None,
        "status": status,
        "remark": _clean(row.get("备注")) or None,
        "is_framework": _clean(row.get("是否框架")).startswith("是"),
        "_type_code": code,
    }
    tag_txt = _clean(row.get("标签(逗号分隔)"))
    payload["tags"] = [t.strip() for t in tag_txt.split(",") if t.strip()] if tag_txt else []
    w_amount = _to_dec(row.get("质保金额"), "质保金额", allow_none=True)
    w_rate = _to_dec(row.get("质保比例%"), "质保比例%", allow_none=True)
    w_start = _to_date(row.get("质保生效日期"))
    w_months = row.get("质保期限(月)")
    if w_amount or w_rate or w_start or w_months not in (None, ""):
        payload["has_warranty"] = True
        payload["warranty_amount"] = float(w_amount) if w_amount is not None else None
        payload["warranty_rate"] = float(w_rate) if w_rate is not None else None
        payload["warranty_start"] = w_start
        try:
            payload["warranty_months"] = int(w_months) if w_months not in (None, "") else None
        except (ValueError, TypeError):
            raise ValueError(f"质保期限(月)应为整数：{w_months}")
    else:
        payload["has_warranty"] = False
    # 主体码（我方公司）
    subjects = get_subjects(db)
    raw_subject = _clean(row.get("主体码"))
    if raw_subject:
        hit = next((s for s in subjects if raw_subject.upper() == s["code"] or raw_subject == s["name"]), None)
        if hit is None:
            raise ValueError(f"未知主体码：{raw_subject}（可用 {', '.join(s['code'] for s in subjects)}）")
        payload["subject_code"] = hit["code"]
    else:
        payload["subject_code"] = subjects[0]["code"] if subjects else ""
        if not payload["subject_code"]:
            raise ValueError("尚未配置我方公司主体，请先在系统设置中添加")
    # 所属框架（按编号查询）
    payload["_parent_no"] = _clean(row.get("所属框架编号"))
    payload["_amount_or_items"] = amount
    return payload


def _parse_items(sheet_rows) -> dict[int, list[dict]]:
    """按『主档序号』分组（兼容旧模板按 合同编号 分组则由旧键处理，这里统一返回序号键 + 编号键）。"""
    grouped: dict[int, list] = {}
    legacy_by_no: dict[str, list] = {}
    for _, row in sheet_rows:
        ord_txt = _clean(row.get("主档序号"))
        no = _clean(row.get("合同编号(对应主档)") or row.get("合同编号"))
        qty = _to_dec(row.get("数量"), "数量", allow_none=True)
        price = _to_dec(row.get("单价"), "单价", allow_none=True)
        name = _clean(row.get("名称"))
        item = {
            "item_type": _clean(row.get("类型")) or "采购",
            "name": name,
            "spec": _clean(row.get("规格型号")),
            "qty": float(qty) if qty is not None else 0,
            "unit_price": float(price) if price is not None else 0,
            "remark": _clean(row.get("备注")) or None,
        }
        if not name and not item["qty"] and not item["unit_price"] and not item["spec"]:
            continue
        if ord_txt.isdigit():
            grouped.setdefault(int(ord_txt), []).append(item)
        elif no:
            legacy_by_no.setdefault(no, []).append(item)
    return {"by_ord": grouped, "by_no": legacy_by_no}


@router.get("/template.xlsx")
def download_template():
    buf = io.BytesIO(build_template_bytes())
    filename = "合同导入模板.xlsx"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=headers)


@router.post("/contracts")
async def import_contracts(file: UploadFile = File(...), db: Session = Depends(get_db)):
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(status_code=422, detail="无法解析文件，请使用系统模板（.xlsx）")

    main_name = next((n for n in wb.sheetnames if n.startswith("合同")), None)
    items_name = next((n for n in wb.sheetnames if n.startswith("行项")), None)
    if main_name is None:
        raise HTTPException(status_code=422, detail="模板缺少『合同(主档)』工作表")

    _, main_rows = _read_rows(wb[main_name])
    item_data = _parse_items(_read_rows(wb[items_name])[1] if items_name else [])
    items_by_ord: dict[int, list] = item_data["by_ord"]
    items_by_no: dict[str, list] = item_data["by_no"]
    if len(main_rows) > LIMIT_MAIN:
        raise HTTPException(status_code=422, detail=f"单次最多导入 {LIMIT_MAIN} 条合同")

    errors: list[dict] = []
    success = 0
    seen_no: set[str] = set()
    seen_ord: set[int] = set()

    for sheet_row, row in main_rows:
        ord_no = _ordinal(row, 0)
        display_no = _clean(row.get("合同编号"))
        try:
            payload = _parse_main(db, row)
            code = payload.pop("_type_code")
            no = payload["contract_no"]

            # 自动编号：留空时按 类型+主体+年（月取签期）
            if not no:
                if code == "OTH":
                    raise ValueError("类型『其他(历史)』不参与自动编号：请手填合同编号或改用其他类型")
                ref = None
                if payload.get("sign_date"):
                    try:
                        ref = date.fromisoformat(payload["sign_date"][:10])
                    except ValueError:
                        ref = None
                no = next_number(db, code, payload["subject_code"], ref)
                payload["contract_no"] = no

            if no in seen_no:
                raise ValueError("文件内合同编号重复")
            if db.query(Contract).filter(Contract.contract_no == no).first():
                raise ValueError("合同编号与库内已有合同重复")

            # 行项：优先按主档序号（新模板），旧模板退回按编号
            items = items_by_ord.get(ord_no) if ord_no else None
            if items is None:
                items = items_by_no.get(no, [])
            if items:
                payload["items"] = items
            else:
                amount = payload.pop("_amount_or_items")
                if amount is None or amount <= 0:
                    raise ValueError("该合同没有行项，金额必须填写（或到『行项明细』挂行）")
                payload["amount"] = float(amount)

            parent_no = payload.pop("_parent_no", "") or ""
            if parent_no:
                parent = db.query(Contract).filter(Contract.contract_no == parent_no).first()
                if parent is None or parent.deleted or not parent.is_framework:
                    raise ValueError(f"所属框架编号不存在或不是框架合同：{parent_no}")
                payload["parent_id"] = parent.id

            if payload["is_framework"] and payload.get("parent_id"):
                raise ValueError("框架合同不能同时挂到其他框架下")
            if payload["status"] not in STATUSES:
                raise ValueError(f"无效状态：{payload['status']}")

            create_contract(payload=payload, db=db)  # 复用校验/行项/标签/框架标签/类型标签逻辑
            seen_no.add(no)
            if ord_no:
                seen_ord.add(ord_no)
            success += 1
        except HTTPException as e:
            errors.append({"row": sheet_row, "contract_no": display_no or (f"第{ord_no}条" if ord_no else ""),
                           "reason": e.detail})
        except ValueError as e:
            errors.append({"row": sheet_row, "contract_no": display_no or (f"第{ord_no}条" if ord_no else ""),
                           "reason": str(e)})

    return {"success": success, "fail": len(errors), "errors": errors}
