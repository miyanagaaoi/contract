"""Excel 台账导出 API（MVP2 需求③：导出项可配置）。

GET /api/export/contracts.xlsx?<筛选参数>&cols=key1,key2,...
- cols 缺省 = 重要列（系统默认打钩的列）
- 复用 _filtered_query（单一数据源），导出的筛选语义与列表一致（AC-11）
"""
from __future__ import annotations

import io
from datetime import date
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..database import get_db
from ..models import Contract
from .contracts import _filtered_query

router = APIRouter(prefix="/api/export", tags=["export"])

# 导出列规格：key → label/是否重要(默认勾)/宽度
COLUMNS: list[dict] = [
    # ---- 重要列（默认勾选导出） ----
    {"key": "contract_no", "label": "合同编号", "important": True, "width": 14},
    {"key": "name", "label": "合同名称", "important": True, "width": 24},
    {"key": "type", "label": "类型", "important": True, "width": 8},
    {"key": "party_a", "label": "甲方", "important": True, "width": 18},
    {"key": "party_b", "label": "乙方", "important": True, "width": 18},
    {"key": "sign_date", "label": "签订日期", "important": True, "width": 12},
    {"key": "subject_matter", "label": "标的物(行项摘要)", "important": True, "width": 24},
    {"key": "amount", "label": "合同金额", "important": True, "width": 12},
    {"key": "paid_amount", "label": "累计已付", "important": True, "width": 12},
    {"key": "payment_ratio", "label": "付款比例%", "important": True, "width": 10},
    {"key": "status", "label": "状态", "important": True, "width": 12},
    {"key": "owner_name", "label": "经办人", "important": True, "width": 10},
    {"key": "tags", "label": "标签", "important": True, "width": 18},
    {"key": "warranty_end", "label": "质保到期日", "important": True, "width": 12},
    # ---- 可选列（默认不勾） ----
    {"key": "currency", "label": "币种", "important": False, "width": 8},
    {"key": "arrival_status", "label": "到货状态", "important": False, "width": 10},
    {"key": "expected_arrival_date", "label": "预计到货日期", "important": False, "width": 13},
    {"key": "is_framework", "label": "是否框架", "important": False, "width": 8},
    {"key": "parent_no", "label": "所属框架编号", "important": False, "width": 16},
    {"key": "warranty_amount", "label": "质保金金额", "important": False, "width": 12},
    {"key": "warranty_rate", "label": "质保金比例%", "important": False, "width": 10},
    {"key": "warranty_start", "label": "质保生效日期", "important": False, "width": 13},
    {"key": "warranty_months", "label": "质保期限(月)", "important": False, "width": 10},
    {"key": "warranty_released", "label": "质保状态", "important": False, "width": 10},
    {"key": "remark", "label": "备注", "important": False, "width": 18},
]

IMPORTANT_KEYS = [c["key"] for c in COLUMNS if c["important"]]


def export_column_meta() -> dict:
    return {"columns": COLUMNS, "default_cols": IMPORTANT_KEYS}


def _value(c: Contract, key: str):
    def _parent_no() -> str:
        return c.parent.contract_no if c.parent_id and c.parent else ""

    fn = {
        "contract_no": lambda: c.contract_no,
        "name": lambda: c.name,
        "type": lambda: c.type,
        "party_a": lambda: c.party_a,
        "party_b": lambda: c.party_b,
        "sign_date": lambda: c.sign_date.isoformat() if c.sign_date else "",
        "subject_matter": lambda: c.subject_matter,
        "amount": lambda: float(c.amount) if c.amount is not None else "",
        "currency": lambda: c.currency,
        "paid_amount": lambda: float(c.paid_amount) if c.paid_amount is not None else "",
        "payment_ratio": lambda: round(float(c.payment_ratio), 2) if c.payment_ratio is not None else "",
        "status": lambda: c.status,
        "owner_name": lambda: c.owner_name or "",
        "tags": lambda: ",".join(t.name for t in c.tags),
        "arrival_status": lambda: c.arrival_status,
        "expected_arrival_date": lambda: c.expected_arrival_date.isoformat() if c.expected_arrival_date else "",
        "is_framework": lambda: "是" if c.is_framework else "否",
        "parent_no": _parent_no,
        "warranty_amount": lambda: float(c.warranty_amount) if c.warranty_amount is not None else "",
        "warranty_rate": lambda: float(c.warranty_rate) if c.warranty_rate is not None else "",
        "warranty_start": lambda: c.warranty_start.isoformat() if c.warranty_start else "",
        "warranty_months": lambda: c.warranty_months or "",
        "warranty_end": lambda: c.warranty_end.isoformat() if c.warranty_end else "",
        "warranty_released": lambda: "已释放" if c.warranty_released else ("未处理" if c.has_warranty else ""),
        "remark": lambda: c.remark or "",
    }
    return fn[key]()


@router.get("/contracts.xlsx")
def export_contracts(
    keyword: str | None = Query(None),
    status: str | None = Query(None),
    contract_type: str | None = Query(None, alias="type"),
    is_framework: bool | None = Query(None),
    include_deleted: bool = Query(False),
    owner: str | None = Query(None),
    tags: str | None = Query(None),
    sign_from: str | None = Query(None),
    sign_to: str | None = Query(None),
    cols: str | None = Query(None, description="逗号分隔的导出列 key；缺省=重要列"),
    db: Session = Depends(get_db),
):
    if cols:
        wanted = [k.strip() for k in cols.split(",") if k.strip()]
        by_key = {c["key"]: c for c in COLUMNS}
        for k in wanted:
            if k not in by_key:
                raise HTTPException(status_code=422, detail=f"未知导出列: {k}")
    else:
        wanted = list(IMPORTANT_KEYS)
    ordered = [c for c in COLUMNS if c["key"] in wanted]

    q = _filtered_query(db, include_deleted=include_deleted, keyword=keyword, owner=owner,
                        status=status, contract_type=contract_type, is_framework=is_framework,
                        sign_from=sign_from, sign_to=sign_to, tags=tags)
    contracts = q.order_by(Contract.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "合同台账"
    ws.append([c["label"] for c in ordered])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for c in contracts:
        ws.append([_value(c, col["key"]) for col in ordered])
    for idx, col in enumerate(ordered, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = col["width"]
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"合同台账_{date.today().isoformat()}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=headers)
