"""生成《合同导入模板样例.xlsx》（MVP2 需求②，供确认导入方案用）。
输出: import_template/合同导入模板样例.xlsx
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path

OUT = Path(__file__).resolve().parents[2] / "import_template" / "合同导入模板样例.xlsx"

HEAD_MAIN = ["合同编号", "合同名称", "类型", "甲方", "乙方", "签订日期", "金额",
             "已付金额", "经办人", "状态", "标签(逗号分隔)", "备注", "质保金额", "质保比例%",
             "质保生效日期", "质保期限(月)"]
HEAD_ITEMS = ["合同编号(对应主档)", "序号", "类型", "名称", "规格型号", "数量", "单价", "备注"]

MAIN_SAMPLE = [
    ["CG-2026-101", "示例：新购设备合同", "采购", "本公司（甲方示例）", "XX 供应商", "2026-09-01", 132000, 0, "张三", "内部审批中", "采购,项目A", "", 6600, 5, "2026-09-01", 12],
    ["XS-2026-201", "示例：产品销售", "销售", "客户 M 公司", "本公司（乙方示例）", "2026-09-05", 6000, 6000, "李四", "到货", "销售", "", 300, 5, "2026-09-05", 12],
]
ITEMS_SAMPLE = [
    ["CG-2026-101", 1, "采购", "X 系列设备", "X-2000", 10, 12000, "含安装调试"],
    ["CG-2026-101", 2, "采购", "配套耗材", "HC-02", 2, 6000, ""],
    ["XS-2026-201", 1, "销售", "X 设备(整机)", "X-2000", 1, 6000, "已收款"],
]


def _style(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for idx, w in enumerate([16, 24, 10, 18, 18, 12, 12, 12, 10, 14, 22, 16, 12, 10, 13, 12], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "合同(主档)"
    _style(ws_main, HEAD_MAIN)
    for row in MAIN_SAMPLE:
        ws_main.append(row)
    ws_items = wb.create_sheet("行项明细")
    _style(ws_items, HEAD_ITEMS)
    for row in ITEMS_SAMPLE:
        ws_items.append(row)
    note = wb.create_sheet("填写说明")
    note.column_dimensions["A"].width = 100
    tips = [
        "填写说明：",
        "1. 两个工作表都要填：『合同(主档)』每行一个合同；『行项明细』每行一个明细，同一合同的明细序号连续。",
        "2. 金额/已付/数量/单价/比例 只填数字；签订日期/质保生效日期填 YYYY-MM-DD。",
        "3. 标签列多个标签用英文逗号分隔（如：采购,项目A），系统里不存在的标签会自动创建。",
        "4. 金额留空 = 由行项合计自动计算；若某合同没有行项（如框架/预估），金额必须手填。",
        "5. 同一合同编号只出现一次；重复、必填缺失、日期或数字格式错误的行会整条跳过并写入错误报告。",
        "6. 状态可选：内部审批中/集团审批中/已签订/付款中/发货/到货/已终止（留空默认为内部审批中）。",
        "7. 示例行可整行删除后再填写，导入前建议用真实数据先试 3~5 条。",
    ]
    for t in tips:
        note.append([t])
    wb.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
    # 自检可读回
    wb = load_workbook(OUT, read_only=True)
    print("sheets:", wb.sheetnames, "main rows:", wb["合同(主档)"].max_row, "items rows:", wb["行项明细"].max_row)
